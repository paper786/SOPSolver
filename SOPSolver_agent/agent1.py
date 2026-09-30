from __future__ import annotations

import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence


@dataclass
class ProblemSpec:
    suite: str
    func_num: int
    dimension: int
    max_iterations: int = 3000
    population_num: int = 100
    time_limit_sec: int = 300
    accept_gap: float = 0.0
    strict_gap: bool = False

    @property
    def function_name(self) -> str:
        return f"{self.suite.upper()}_F{self.func_num}"

    @property
    def file_stem(self) -> str:
        return f"{self.suite.lower()}_fun{self.func_num}_{self.dimension}D"

    def candidate_stem(self, agent_id: str) -> str:
        return f"{self.file_stem}_{agent_id}"


@dataclass
class SolverConfig:
    project_root: Path
    api_key: str
    api_base_url: str = "https://api.openai.com/v1"
    model: str = "gpt-5.5"
    temperature: float = 0.75
    request_timeout_sec: int = 120
    max_tokens: int = 8192
    literature_char_limit: int = 180000
    prompt_char_limit: int = 24000
    feedback_char_limit: int = 16000

    @property
    def compare_code_dir(self) -> Path:
        return self.project_root / "compare_code"

    @property
    def final_code_dir(self) -> Path:
        return self.project_root / "final_code"

    @property
    def evaluate_prompt_dir(self) -> Path:
        return self.project_root / "prompt_agent_evaluate_md"

    @property
    def artifacts_dir(self) -> Path:
        return self.project_root / "sopsolver_artifacts"

    @property
    def statistics_dir(self) -> Path:
        return self.project_root / "Statistics"

    @property
    def literature_dir(self) -> Path:
        return self.project_root / "Literature"


@dataclass
class GeneratedAlgorithm:
    agent_id: str
    round_index: int
    path: Path
    code: str
    raw_response: str
    algorithm_name: str = ""
    combination_number: Optional[int] = None
    created_at: str = ""


class APIError(RuntimeError):
    pass


class NoCodeBlockError(RuntimeError):
    pass


class OpenAICompatibleClient:
    """Small OpenAI-compatible chat client implemented with the standard library."""

    def __init__(
        self,
        api_key: str,
        base_url: str = "https://api.openai.com/v1",
        model: str = "gpt-5.5",
        timeout_sec: int = 120,
    ) -> None:
        self.api_key = api_key.strip()
        self.endpoint = self._normalize_endpoint(base_url)
        self.model = model
        self.timeout_sec = timeout_sec

    @staticmethod
    def _normalize_endpoint(base_url: str) -> str:
        base = (base_url or "https://api.openai.com/v1").strip().rstrip("/")
        if base.endswith("/chat/completions"):
            return base
        return f"{base}/chat/completions"

    def chat(
        self,
        messages: Sequence[Dict[str, str]],
        *,
        temperature: float = 0.75,
        max_tokens: int = 8192,
    ) -> str:
        if not self.api_key:
            raise APIError("API key is empty. Set SOPSOLVER_API_KEY or pass --api-key.")

        payload = {
            "model": self.model,
            "messages": list(messages),
            "temperature": temperature,
            "max_tokens": max_tokens,
        }
        request = urllib.request.Request(
            self.endpoint,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=self.timeout_sec) as response:
                data = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            body = exc.read().decode("utf-8", errors="replace")
            raise APIError(f"API HTTP {exc.code}: {body}") from exc
        except Exception as exc:  # pragma: no cover - network-dependent.
            raise APIError(f"API request failed: {exc}") from exc

        try:
            return data["choices"][0]["message"]["content"]
        except Exception as exc:
            raise APIError(f"Unexpected API response: {data}") from exc


def ensure_product_directories(config: SolverConfig) -> None:
    for path in [
        config.compare_code_dir,
        config.final_code_dir,
        config.evaluate_prompt_dir,
        config.artifacts_dir,
        config.statistics_dir,
    ]:
        path.mkdir(parents=True, exist_ok=True)


def read_text(path: Path, *, max_chars: Optional[int] = None) -> str:
    encodings = ("utf-8-sig", "utf-8", "gb18030", "cp936")
    last_error: Optional[Exception] = None
    for encoding in encodings:
        try:
            text = path.read_text(encoding=encoding)
            if max_chars is not None and len(text) > max_chars:
                return text[:max_chars] + "\n\n[TRUNCATED]\n"
            return text
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Cannot read {path}: {last_error}")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def extract_code_block(message: str, preferred_language: str = "matlab") -> str:
    blocks = re.findall(r"```([A-Za-z0-9_+-]*)\s*\n(.*?)```", message, re.DOTALL)
    if not blocks:
        raise NoCodeBlockError("The API response did not contain a fenced code block.")

    preferred = [
        code.strip()
        for lang, code in blocks
        if lang.strip().lower() in {preferred_language.lower(), "m"}
    ]
    if preferred:
        return max(preferred, key=len)
    return max((code.strip() for _, code in blocks), key=len)


def extract_first_number(patterns: Iterable[str], text: str) -> Optional[int]:
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            try:
                return int(float(match.group(1)))
            except ValueError:
                return None
    return None


def normalize_matlab_function_name(code: str, target_name: str) -> str:
    pattern = r"(?im)^(\s*function\s+(?:\[[^\]]+\]\s*=\s*|[A-Za-z]\w*\s*=\s*)?)([A-Za-z]\w*)(\s*\()"
    replacement = rf"\g<1>{target_name}\g<3>"
    updated, count = re.subn(pattern, replacement, code, count=1)
    if count == 0:
        raise NoCodeBlockError("The MATLAB code must start with a function definition.")
    return updated


def parse_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text == "-":
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def load_openpyxl():
    try:
        import openpyxl  # type: ignore

        return openpyxl
    except Exception as exc:
        fallback = (
            Path.home()
            / ".cache"
            / "codex-runtimes"
            / "codex-primary-runtime"
            / "dependencies"
            / "python"
            / "Lib"
            / "site-packages"
        )
        if fallback.exists() and str(fallback) not in sys.path:
            sys.path.insert(0, str(fallback))
            try:
                import openpyxl  # type: ignore

                return openpyxl
            except Exception:
                pass
        raise RuntimeError(
            "openpyxl is required for Excel reading/writing. Install it with "
            "`pip install openpyxl`, or run this product with the Codex bundled Python."
        ) from exc


def find_best_record(config: SolverConfig, problem: ProblemSpec) -> Dict[str, Any]:
    workbook = config.statistics_dir / "BestRecord.xlsx"
    if not workbook.exists():
        return {"available": False, "best": None, "source": str(workbook)}

    openpyxl = load_openpyxl()
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb.active
    target_function = problem.function_name.upper()
    target_dimension = str(problem.dimension)

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if str(ws.cell(row, col).value).strip().lower() != "function":
                continue
            headers: Dict[str, int] = {}
            for c in range(col, min(ws.max_column, col + 12) + 1):
                value = ws.cell(row, c).value
                if value is not None:
                    headers[str(value).strip().lower()] = c
            if "dimension" not in headers or "best" not in headers:
                continue
            func_col = col
            dim_col = headers["dimension"]
            best_col = headers["best"]
            for r in range(row + 1, ws.max_row + 1):
                func_value = ws.cell(r, func_col).value
                dim_value = ws.cell(r, dim_col).value
                if str(func_value).strip().upper() == target_function and str(dim_value).strip() == target_dimension:
                    return {
                        "available": True,
                        "best": parse_float(ws.cell(r, best_col).value),
                        "row": r,
                        "column": best_col,
                        "algorithm_abbr": ws.cell(r, headers.get("algorithm_abbr", best_col)).value,
                        "algorithm_full_name": ws.cell(r, headers.get("algorithm_full_name", best_col)).value,
                        "source": str(workbook),
                    }
    return {"available": False, "best": None, "source": str(workbook)}


def update_statistics_workbook(
    workbook_path: Path,
    problem: ProblemSpec,
    stats: Dict[str, Any],
) -> None:
    openpyxl = load_openpyxl()
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    if workbook_path.exists():
        wb = openpyxl.load_workbook(workbook_path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(
            [
                "Function",
                "Dimension",
                "Best",
                "Ave",
                "Std",
                "Worst",
                "Time",
                "Iteration",
                "Population_num",
                "Algorithm_combination",
                "combination_number",
            ]
        )

    headers = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)}
    function_col = headers.get("Function")
    dimension_col = headers.get("Dimension")
    if not function_col or not dimension_col:
        raise RuntimeError(f"{workbook_path} does not contain Function/Dimension headers.")

    target_row = None
    for r in range(2, ws.max_row + 1):
        if (
            str(ws.cell(r, function_col).value).strip().upper() == problem.function_name.upper()
            and str(ws.cell(r, dimension_col).value).strip() == str(problem.dimension)
        ):
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
        ws.cell(target_row, function_col).value = problem.function_name
        ws.cell(target_row, dimension_col).value = problem.dimension

    mapping = {
        "Best": "best",
        "Ave": "ave",
        "Std": "std",
        "Worst": "worst",
        "Time": "time",
        "Iteration": "iteration",
        "Population_num": "population_num",
        "Algorithm_combination": "algorithm_combination",
        "combination_number": "combination_number",
    }
    for header, key in mapping.items():
        if header in headers and key in stats:
            ws.cell(target_row, headers[header]).value = stats[key]
    wb.save(workbook_path)


def collect_literature_context(config: SolverConfig, problem: ProblemSpec) -> str:
    lines: List[str] = []
    info_path = config.statistics_dir / "Algorithm_information.xlsx"
    if info_path.exists():
        try:
            openpyxl = load_openpyxl()
            wb = openpyxl.load_workbook(info_path, data_only=True)
            ws = wb.active
            headers = {
                str(ws.cell(1, c).value).strip(): c
                for c in range(1, ws.max_column + 1)
                if ws.cell(1, c).value is not None
            }
            for r in range(2, ws.max_row + 1):
                abbr = ws.cell(r, headers.get("Algorithm_Abbr", 1)).value
                full = ws.cell(r, headers.get("Algorithm_Full_Name", 2)).value
                year = ws.cell(r, headers.get("Year", 3)).value
                cec = ws.cell(r, headers.get("CEC_Benchmark", 4)).value
                funcs = ws.cell(r, headers.get("Function_Number", 5)).value
                if not abbr or not full:
                    continue
                cec_text = str(cec or "")
                score = 0
                if problem.suite.upper() in cec_text.upper():
                    score += 2
                if str(problem.func_num) in str(funcs or ""):
                    score += 1
                prefix = "*" if score else "-"
                lines.append(f"{prefix} {abbr}: {full} ({year}); CEC={cec}; Function={funcs}")
        except Exception as exc:
            lines.append(f"[Algorithm_information.xlsx could not be read: {exc}]")

    if config.literature_dir.exists():
        pdf_names = sorted(p.name for p in config.literature_dir.glob("*.pdf"))
        lines.append("\nLiterature PDF files available under Literature:")
        lines.extend(f"- {name}" for name in pdf_names[:160])

    context = "\n".join(lines)
    if len(context) > config.literature_char_limit:
        return context[: config.literature_char_limit] + "\n\n[TRUNCATED]\n"
    return context


class SOPDesignAgent:
    agent_id = "Agent1"
    role_hint = "Design the first independent candidate algorithm."
    difference_hint = (
        "Prefer a balanced metaheuristic built from literature-inspired operators. "
        "Do not copy any previous code."
    )

    def __init__(self, config: SolverConfig) -> None:
        self.config = config
        self.client = OpenAICompatibleClient(
            config.api_key,
            config.api_base_url,
            config.model,
            config.request_timeout_sec,
        )

    def generate_algorithm(
        self,
        problem: ProblemSpec,
        *,
        prompt_text: str,
        call_rules_text: str,
        best_record: Dict[str, Any],
        feedback_text: str = "",
        peer_summary: str = "",
        previous_code: str = "",
        round_index: int = 1,
    ) -> GeneratedAlgorithm:
        ensure_product_directories(self.config)
        literature_context = collect_literature_context(self.config, problem)
        target_stem = problem.candidate_stem(self.agent_id)
        messages = self._build_messages(
            problem,
            target_stem,
            prompt_text,
            call_rules_text,
            literature_context,
            best_record,
            feedback_text,
            peer_summary,
            previous_code,
            round_index,
        )
        response = self.client.chat(
            messages,
            temperature=self.config.temperature,
            max_tokens=self.config.max_tokens,
        )
        save_generation_artifacts(self.config, problem, self.agent_id, round_index, messages, response)
        code = extract_code_block(response, "matlab")
        code = normalize_matlab_function_name(code, target_stem)
        path = self.config.compare_code_dir / f"{target_stem}.m"
        write_text(path, code)
        algorithm_name = self._extract_algorithm_name(response, code)
        combination_number = extract_first_number(
            [r"combination[_\s-]*number\s*[:=]\s*(\d+)", r"strategy\s*number\s*[:=]\s*(\d+)"],
            response + "\n" + code,
        )
        return GeneratedAlgorithm(
            agent_id=self.agent_id,
            round_index=round_index,
            path=path,
            code=code,
            raw_response=response,
            algorithm_name=algorithm_name,
            combination_number=combination_number,
            created_at=time.strftime("%Y-%m-%d %H:%M:%S"),
        )

    def _build_messages(
        self,
        problem: ProblemSpec,
        target_stem: str,
        prompt_text: str,
        call_rules_text: str,
        literature_context: str,
        best_record: Dict[str, Any],
        feedback_text: str,
        peer_summary: str,
        previous_code: str,
        round_index: int,
    ) -> List[Dict[str, str]]:
        base_prompt = prompt_text[: self.config.prompt_char_limit]
        feedback = feedback_text[: self.config.feedback_char_limit]
        previous = previous_code[-10000:] if previous_code else ""
        best_value = best_record.get("best")
        system = (
            "You are one agent in SOPSolver, a multi-agent framework for single-objective "
            "continuous optimization. Generate honest, runnable MATLAB metaheuristic code. "
            "Never fabricate results, never hard-code known optima, never read shift/rotation "
            "data directly, and never modify benchmark functions."
        )
        user = f"""
Current role: {self.agent_id}. {self.role_hint}
Agent-specific design requirement: {self.difference_hint}

Problem:
- Suite: {problem.suite}
- Function number: F{problem.func_num}
- Dimension: {problem.dimension}
- MATLAB function name and file stem: {target_stem}
- Max iterations for generated code default: {problem.max_iterations}
- Population size default: {problem.population_num}
- Single-run time limit hint: {problem.time_limit_sec} seconds
- BestRecord reference value: {best_value}
- Acceptance gap in this Python framework: {problem.accept_gap}; strict gap: {problem.strict_gap}

Main SOPSolver prompt excerpt:
{base_prompt}

Benchmark call rules excerpt:
{call_rules_text[:12000]}

Literature and algorithm inventory:
{literature_context}

Evaluate_Agent feedback for this problem, if any:
{feedback or "[No previous feedback for this problem.]"}

Peer/previous-agent summary, if any:
{peer_summary or "[No peer summary.]"}

Previous code from this same agent, if any:
{previous or "[No previous code.]"}

Mandatory MATLAB interface:
```matlab
function result = {target_stem}(seed, options)
% Candidate {self.agent_id} for {problem.function_name} {problem.dimension}D.
% The function must return a struct named result.
end
```

Implementation requirements:
1. Use only MATLAB code in one complete fenced matlab code block.
2. The primary function name must be exactly {target_stem}.
3. Add paths inside the function using:
   root_dir = fileparts(fileparts(mfilename('fullpath')));
   addpath(fullfile(root_dir, '{problem.suite}'));
4. For CEC2017 use objective calls such as CEC2017_F01(x), CEC2017_F03(x), etc.; for other suites follow the local call rules.
5. Use bound repair/projection. Default CEC2014/CEC2017 bounds are [-100, 100].
6. Return at least:
   result.best_value
   result.record_value
   result.best_position
   result.convergence_curve
   result.runtime
   result.iteration
   result.population_num
   result.algorithm_combination
   result.combination_number
   result.agent_id
7. Include result.record_value as the minimization value used for comparison.
8. Use English comments and fprintf text.
9. Do not output placeholder code. Do not call non-existent helper functions unless you define them in the same file.
10. Round {round_index}: if feedback shows poor progress, redesign operators instead of only tuning parameters.
11. All candidate solutions must be represented as 1-by-dim row vectors.
12. The population matrix must be pop_size-by-dim.
13. When selecting one solution from the population or archive, always use row indexing, e.g., x = pop(idx, :).
14. Before arithmetic operations, reshape vectors by x = reshape(x, 1, dim).
15. Use element-wise operators .*, ./, and .^ for vector operations.
16. Never add or subtract a full population matrix from a single solution vector unless dimensions are explicitly matched.
17. Before returning code, check that every generated new_solution is a 1-by-dim row vector.

Response format:
# Algorithm
Short explanation of selected literature-inspired algorithm or fusion strategy.
# MATLAB Code
```matlab
<complete MATLAB code>
```
"""
        return [{"role": "system", "content": system}, {"role": "user", "content": user}]

    @staticmethod
    def _extract_algorithm_name(response: str, code: str) -> str:
        for pattern in [
            r"#\s*Algorithm\s*\n\s*([^\n]+)",
            r"algorithm[_\s-]*combination\s*=\s*['\"]([^'\"]+)['\"]",
            r"result\.algorithm_combination\s*=\s*['\"]([^'\"]+)['\"]",
        ]:
            match = re.search(pattern, response + "\n" + code, re.IGNORECASE)
            if match:
                return match.group(1).strip()
        return ""


class Agent1(SOPDesignAgent):
    agent_id = "Agent1"
    role_hint = (
        "Design the first candidate independently. Favor robust global search, "
        "archive use, adaptive perturbation, and problem-specific balance."
    )
    difference_hint = (
        "Start from one or several literature metaheuristics with a clear strategy "
        "number. Avoid simply imitating Agent2; prefer exploration-first designs."
    )


class OfflineTemplateAgent(SOPDesignAgent):
    """Deterministic fallback agent for pipeline verification without an API key."""

    def __init__(self, config: SolverConfig, agent_id: str = "Agent1") -> None:
        super().__init__(config)
        self.agent_id = agent_id

    def generate_algorithm(
        self,
        problem: ProblemSpec,
        *,
        prompt_text: str,
        call_rules_text: str,
        best_record: Dict[str, Any],
        feedback_text: str = "",
        peer_summary: str = "",
        previous_code: str = "",
        round_index: int = 1,
    ) -> GeneratedAlgorithm:
        ensure_product_directories(self.config)
        target_stem = problem.candidate_stem(self.agent_id)
        code = build_offline_template_code(problem, target_stem, self.agent_id)
        path = self.config.compare_code_dir / f"{target_stem}.m"
        write_text(path, code)
        response = (
            "# Offline Template\n"
            "No API key was used. This deterministic baseline is intended to "
            "verify the SOPSolver orchestration, MATLAB execution, result "
            "recording, and artifact-retention pipeline.\n"
        )
        save_generation_artifacts(
            self.config,
            problem,
            self.agent_id,
            round_index,
            [
                {
                    "role": "system",
                    "content": "Offline deterministic SOPSolver template generation.",
                },
                {
                    "role": "user",
                    "content": (
                        f"Generate {target_stem} for {problem.function_name} "
                        f"{problem.dimension}D using the local template."
                    ),
                },
            ],
            response + "\n```matlab\n" + code + "\n```",
        )
        return GeneratedAlgorithm(
            agent_id=self.agent_id,
            round_index=round_index,
            path=path,
            code=code,
            raw_response=response,
            algorithm_name="Offline deterministic baseline",
            combination_number=0,
            created_at=time.strftime("%Y-%m-%d %H:%M:%S"),
        )


def save_generation_artifacts(
    config: SolverConfig,
    problem: ProblemSpec,
    agent_id: str,
    round_index: int,
    messages: Sequence[Dict[str, str]],
    response: str,
) -> None:
    artifact_dir = config.artifacts_dir / problem.file_stem / f"round_{round_index:03d}" / agent_id
    artifact_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "suite": problem.suite,
        "function": problem.func_num,
        "dimension": problem.dimension,
        "agent_id": agent_id,
        "round_index": round_index,
        "model": config.model,
        "api_base_url": config.api_base_url,
        "messages": list(messages),
        "response": response,
    }
    write_text(artifact_dir / "llm_request_response.json", json.dumps(payload, ensure_ascii=False, indent=2))
    prompt_md = [
        f"# {problem.file_stem} {agent_id} Round {round_index}",
        "",
        "## Request",
        "",
    ]
    for message in messages:
        prompt_md.extend([f"### {message.get('role', 'unknown')}", "", message.get("content", ""), ""])
    prompt_md.extend(["## Response", "", response])
    write_text(artifact_dir / "llm_request_response.md", "\n".join(prompt_md))


def build_offline_template_code(problem: ProblemSpec, target_stem: str, agent_id: str) -> str:
    func_file = f"{problem.suite.upper()}_F{problem.func_num:02d}"
    return f"""function result = {target_stem}(seed, options)
% Offline deterministic SOPSolver baseline for pipeline verification.
% This code is intentionally simple; it proves reproducible execution and
% record keeping when no LLM API key is available.
if nargin < 1 || isempty(seed), seed = 1; end
if nargin < 2 || isempty(options), options = struct(); end
rng(seed, 'twister');
root_dir = fileparts(fileparts(mfilename('fullpath')));
addpath(fullfile(root_dir, '{problem.suite.upper()}'));
dim = {int(problem.dimension)};
lb = -100 * ones(1, dim);
ub = 100 * ones(1, dim);
max_iter = get_option(options, 'max_iteration', {int(problem.max_iterations)});
pop_size = get_option(options, 'population_num', {int(problem.population_num)});
max_runtime = get_option(options, 'max_runtime_sec', {int(problem.time_limit_sec)});
tic;
pop = lb + rand(pop_size, dim) .* (ub - lb);
values = zeros(pop_size, 1);
for i = 1:pop_size
    values(i) = objective(pop(i, :));
end
[best_value, idx] = min(values);
best_position = pop(idx, :);
curve = zeros(max_iter, 1);
sigma = 0.25 * (ub - lb);
for iter = 1:max_iter
    if toc > max_runtime
        curve = curve(1:iter-1);
        break;
    end
    for i = 1:pop_size
        trial = best_position + sigma .* randn(1, dim);
        if rand < 0.35
            peer = pop(randi(pop_size), :);
            trial = trial + 0.4 * rand(1, dim) .* (peer - pop(i, :));
        end
        trial = min(max(trial, lb), ub);
        trial_value = objective(trial);
        if trial_value < values(i)
            pop(i, :) = trial;
            values(i) = trial_value;
            if trial_value < best_value
                best_value = trial_value;
                best_position = trial;
            end
        end
    end
    sigma = max(1e-8, 0.995) .* sigma;
    curve(iter) = best_value;
end
if isempty(curve), curve = best_value; end
result = struct();
result.best_value = best_value;
result.record_value = best_value;
result.best_position = best_position;
result.convergence_curve = curve(:);
result.runtime = toc;
result.iteration = numel(curve);
result.population_num = pop_size;
result.algorithm_combination = 'Offline deterministic stochastic hill-climbing baseline';
result.combination_number = 0;
result.agent_id = '{agent_id}';

function value = objective(x)
value = {func_file}(x(:));
end

function value = get_option(opts, name, default_value)
if isstruct(opts) && isfield(opts, name) && ~isempty(opts.(name))
    value = opts.(name);
else
    value = default_value;
end
end
end
"""
